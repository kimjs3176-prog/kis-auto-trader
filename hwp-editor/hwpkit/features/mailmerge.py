"""메일머지 — 표 형태 자료로 양식을 여러 부 만든다.

원본 `메일머지창` / `메머체워넣기` 는 필드 이름이 '1'~'20' 으로 고정이라 항목이
20개를 넘을 수 없었고, 자료를 사람이 한/글 창에 직접 붙여 넣어야 했다.

여기서는
* CSV / TSV / XLSX 를 직접 읽고(엑셀이 깔려 있지 않아도 된다),
* 첫 줄을 항목 이름으로 써서 이름이 몇 개든 상관없고,
* 파일 이름 규칙(`{이름}_{날짜}` 같은 서식)을 지정할 수 있으며,
* HWP·PDF 로 한 번에 저장한다.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from ..core.document import 문서
from ..core.errors import 입력오류
from . import forms

__all__ = ["자료읽기", "머지결과", "메일머지", "파일이름만들기"]

_금지문자 = re.compile(r'[\\/:*?"<>|\r\n\t]')


def 자료읽기(경로: str | Path) -> list[dict[str, str]]:
    """CSV / TSV / XLSX 를 읽어 [{항목: 값}] 목록으로 만든다."""
    파일 = Path(경로)
    if not 파일.is_file():
        raise 입력오류(f"자료 파일이 없습니다: {파일}")

    확장자 = 파일.suffix.lower()
    if 확장자 in (".csv", ".tsv", ".txt"):
        return _표읽기_텍스트(파일)
    if 확장자 in (".xlsx", ".xlsm"):
        return _표읽기_엑셀(파일)
    raise 입력오류(
        f"읽을 수 없는 자료 형식입니다: {확장자}",
        "CSV, TSV, XLSX 파일을 쓰세요. (엑셀에서 '다른 이름으로 저장 → CSV')",
    )


def _표읽기_텍스트(파일: Path) -> list[dict[str, str]]:
    원문 = None
    for 인코딩 in ("utf-8-sig", "cp949", "utf-8"):
        try:
            원문 = 파일.read_text(encoding=인코딩)
            break
        except UnicodeDecodeError:
            continue
    if 원문 is None:
        raise 입력오류(f"글자 인코딩을 알 수 없습니다: {파일.name}")

    구분 = "\t" if 파일.suffix.lower() == ".tsv" else ","
    읽기 = csv.DictReader(io.StringIO(원문), delimiter=구분)
    결과 = []
    for 줄 in 읽기:
        정리 = {
            (이름 or "").strip(): ("" if 값 is None else str(값).strip())
            for 이름, 값 in 줄.items()
            if (이름 or "").strip()
        }
        if any(정리.values()):
            결과.append(정리)
    return 결과


def _표읽기_엑셀(파일: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as 오류:
        raise 입력오류(
            "엑셀 파일을 읽으려면 openpyxl 이 필요합니다.",
            "`pip install openpyxl` 을 실행하거나 CSV 로 저장해 쓰세요.",
        ) from 오류

    책 = load_workbook(filename=str(파일), data_only=True, read_only=True)
    장 = 책.active
    줄들 = 장.iter_rows(values_only=True)
    머리 = None
    결과: list[dict[str, str]] = []
    for 줄 in 줄들:
        값들 = ["" if 칸 is None else str(칸).strip() for 칸 in 줄]
        if 머리 is None:
            if any(값들):
                머리 = 값들
            continue
        if not any(값들):
            continue
        결과.append(
            {
                이름: (값들[번호] if 번호 < len(값들) else "")
                for 번호, 이름 in enumerate(머리)
                if 이름
            }
        )
    책.close()
    return 결과


def 파일이름만들기(서식: str, 자료: dict[str, str], 순번: int) -> str:
    """`{성명}_{부서}` 같은 서식으로 파일 이름을 만든다."""
    이름 = 서식 or "{순번}"
    이름 = 이름.replace("{순번}", f"{순번:03d}")
    for 항목, 값 in 자료.items():
        이름 = 이름.replace(f"{{{항목}}}", str(값))
    이름 = _금지문자.sub("_", 이름).strip() or f"{순번:03d}"
    return 이름[:100]


@dataclass
class 머지결과:
    """메일머지 실행 결과."""

    전체: int = 0
    성공: list[str] = field(default_factory=list)
    실패: list[tuple[int, str]] = field(default_factory=list)

    @property
    def 문구(self) -> str:
        말 = f"{self.전체}건 중 {len(self.성공)}건 저장"
        if self.실패:
            말 += f", {len(self.실패)}건 실패"
        return 말


def 메일머지(
    문서객체: 문서,
    양식이름: str,
    자료: Iterable[dict[str, str]] | str | Path,
    저장폴더: str | Path,
    파일이름서식: str = "{순번}",
    형식: str = "HWP",
    진행알림: Callable[[int, int, str], None] | None = None,
) -> 머지결과:
    """양식에 자료를 채워 한 건씩 저장한다.

    형식: HWP | HWPX | PDF | HWP+PDF
    """
    줄들 = list(자료읽기(자료) if isinstance(자료, (str, Path)) else 자료)
    if not 줄들:
        raise 입력오류("채울 자료가 없습니다.")

    폴더 = Path(저장폴더)
    폴더.mkdir(parents=True, exist_ok=True)

    양식파일 = forms.양식찾기(양식이름).경로
    결과 = 머지결과(전체=len(줄들))
    형식들 = ["HWP", "PDF"] if 형식.upper() == "HWP+PDF" else [형식.upper()]

    보이기원래 = True
    문서객체.연결.보이기(False)  # 숨겨 두면 훨씬 빠르다
    try:
        for 순번, 줄 in enumerate(줄들, start=1):
            이름 = 파일이름만들기(파일이름서식, 줄, 순번)
            if 진행알림:
                진행알림(순번, len(줄들), 이름)
            try:
                문서객체.열기(str(양식파일))
                forms.양식채우기(문서객체, 줄)
                for 하나 in 형식들:
                    확장자 = {"HWP": ".hwp", "HWPX": ".hwpx", "PDF": ".pdf"}.get(하나, ".hwp")
                    문서객체.다른이름저장(str(폴더 / f"{이름}{확장자}"), 하나)
                결과.성공.append(이름)
            except Exception as 오류:  # noqa: BLE001 - 한 건 실패가 전체를 멈추지 않게
                결과.실패.append((순번, f"{이름}: {오류}"))
            finally:
                문서객체.문서닫기(저장=False)
    finally:
        문서객체.연결.보이기(보이기원래)
    return 결과


def 열려있는문서채우기(문서객체: 문서, 자료: dict[str, str]) -> int:
    """지금 열려 있는 문서에 값 한 벌만 채운다. (원본 `메머체워넣기` 자리)"""
    return forms.양식채우기(문서객체, 자료)
