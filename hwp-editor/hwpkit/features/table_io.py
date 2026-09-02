"""표 ↔ 엑셀/CSV 주고받기.

원본에는 별도 `엑셀라이브러리` 가 있었지만 **엑셀이 깔려 있어야** 하고(COM 자동화),
현재 열려 있는 엑셀 선택 영역만 다룰 수 있었다. 여기서는 파일을 직접 읽고 쓰므로
엑셀 없이도 되고, 어떤 파일이든 지정해 쓸 수 있다.

* 표 → CSV/XLSX 내보내기 (CSV 는 엑셀에서 한글이 깨지지 않게 BOM 을 붙인다)
* CSV/XLSX → 표 가져오기 (머리글 스타일까지 적용)
"""

from __future__ import annotations

import csv
from pathlib import Path

from ..core.document import 문서
from ..core.errors import 상태오류, 입력오류

__all__ = ["격자읽기", "격자쓰기", "표내보내기", "표가져오기"]

_읽을수있는확장자 = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")


def 격자읽기(경로: str | Path, 시트: str | None = None) -> list[list[str]]:
    """CSV/TSV/XLSX 를 행×열 표로 읽는다. (머리글도 한 행으로 그대로 읽는다)"""
    파일 = Path(경로)
    if not 파일.is_file():
        raise 입력오류(f"파일이 없습니다: {파일}")
    확장자 = 파일.suffix.lower()
    if 확장자 not in _읽을수있는확장자:
        raise 입력오류(
            f"읽을 수 없는 형식입니다: {확장자}",
            f"쓸 수 있는 형식: {', '.join(_읽을수있는확장자)}",
        )

    if 확장자 in (".xlsx", ".xlsm"):
        return _엑셀격자(파일, 시트)

    원문 = None
    for 인코딩 in ("utf-8-sig", "cp949", "utf-8"):
        try:
            원문 = 파일.read_text(encoding=인코딩)
            break
        except UnicodeDecodeError:
            continue
    if 원문 is None:
        raise 입력오류(f"글자 인코딩을 알 수 없습니다: {파일.name}")

    구분 = "\t" if 확장자 == ".tsv" else ","
    행들 = [
        [칸.strip() for 칸 in 줄]
        for 줄 in csv.reader(원문.splitlines(), delimiter=구분)
        if any(칸.strip() for 칸 in 줄)
    ]
    if not 행들:
        raise 입력오류(f"내용이 없습니다: {파일.name}")
    return 행들


def _엑셀격자(파일: Path, 시트: str | None) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as 오류:
        raise 입력오류(
            "엑셀 파일을 읽으려면 openpyxl 이 필요합니다.",
            "`pip install openpyxl` 을 실행하거나 CSV 로 저장해 쓰세요.",
        ) from 오류

    책 = load_workbook(filename=str(파일), data_only=True, read_only=True)
    try:
        장 = 책[시트] if 시트 else 책.active
    except KeyError as 오류:
        raise 입력오류(
            f"'{시트}' 시트가 없습니다.", f"있는 시트: {', '.join(책.sheetnames)}"
        ) from 오류

    행들: list[list[str]] = []
    for 줄 in 장.iter_rows(values_only=True):
        값들 = ["" if 칸 is None else str(칸).strip() for 칸 in 줄]
        if any(값들):
            행들.append(값들)
    책.close()
    if not 행들:
        raise 입력오류(f"내용이 없습니다: {파일.name}")
    return 행들


def 격자쓰기(행들: list[list[str]], 경로: str | Path, 형식: str = "CSV") -> Path:
    """행×열 표를 CSV 또는 XLSX 파일로 쓴다."""
    if not 행들:
        raise 입력오류("저장할 내용이 없습니다.")
    파일 = Path(경로)
    형식 = 형식.upper()
    if 형식 == "CSV" and 파일.suffix.lower() != ".csv":
        파일 = 파일.with_suffix(".csv")
    if 형식 == "XLSX" and 파일.suffix.lower() != ".xlsx":
        파일 = 파일.with_suffix(".xlsx")
    파일.parent.mkdir(parents=True, exist_ok=True)

    if 형식 == "CSV":
        # 엑셀에서 한글이 깨지지 않도록 BOM 을 붙인다.
        with 파일.open("w", encoding="utf-8-sig", newline="") as 손잡이:
            csv.writer(손잡이).writerows(행들)
        return 파일

    if 형식 != "XLSX":
        raise 입력오류(f"모르는 저장 형식: {형식}", "CSV 또는 XLSX 를 쓰세요.")

    try:
        from openpyxl import Workbook
    except ImportError as 오류:
        raise 입력오류(
            "XLSX 로 저장하려면 openpyxl 이 필요합니다.",
            "`pip install openpyxl` 을 실행하거나 CSV 로 저장하세요.",
        ) from 오류

    책 = Workbook()
    장 = 책.active
    장.title = "표"
    for 행 in 행들:
        장.append(행)
    책.save(str(파일))
    return 파일


def 표내보내기(
    문서객체: 문서, 경로: str | Path, 형식: str = "CSV"
) -> str:
    """선택한 표(셀 블록)를 CSV/XLSX 로 내보낸다."""
    from .table_calc import 표값읽기

    행들, 범위 = 표값읽기(문서객체)
    if not any(any(칸 for 칸 in 행) for 행 in 행들):
        raise 상태오류("내보낼 내용이 없습니다.", "표에서 셀을 선택한 뒤 다시 누르세요.")
    파일 = 격자쓰기(행들, 경로, 형식)
    return f"{범위.행수}행 {범위.열수}열을 {파일.name} 으로 저장했습니다.\n{파일}"


def 표가져오기(
    문서객체: 문서,
    경로: str | Path,
    시트: str | None = None,
    스타일: str = "기본",
    머리글: bool = True,
    최대행: int | None = None,
) -> str:
    """CSV/XLSX 를 읽어 현재 위치에 표로 넣는다."""
    from . import table_ops

    행들 = 격자읽기(경로, 시트)
    if 최대행:
        행들 = 행들[:최대행]
    열수 = max(len(행) for 행 in 행들)

    with 문서객체.한번에("표 가져오기"), 문서객체.화면정지():
        table_ops.표만들기(문서객체, 열수=열수, 행수=len(행들))
        문서객체.이동하기("표처음셀")
        for 행번호, 행 in enumerate(행들):
            for 열번호 in range(열수):
                문서객체.셀텍스트쓰기(행[열번호] if 열번호 < len(행) else "")
                마지막 = 행번호 == len(행들) - 1 and 열번호 == 열수 - 1
                if not 마지막:
                    문서객체.실행("TableRightCell")
        if 스타일:
            table_ops.표스타일적용(문서객체, 스타일, 머리글있음=머리글)
        문서객체.취소()
    return f"{len(행들)}행 {열수}열 표로 가져왔습니다."
